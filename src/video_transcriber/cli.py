"""CLI interface for video-transcriber."""

import argparse
import sys

from video_transcriber.downloader import download_audio
from video_transcriber.postprocess import clean_segments
from video_transcriber.transcriber import transcribe_audio


def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        prog="video-transcriber",
        description="Transcribe European Parliament webstreaming videos.",
    )
    parser.add_argument(
        "url",
        help=(
            "EP webstreaming URL, e.g. "
            "https://multimedia.europarl.europa.eu/en/webstreaming/"
            "committees_20260317-1430-COMMITTEE-EMPL"
        ),
    )
    parser.add_argument(
        "-o",
        "--output",
        help="Output file path for the transcript (default: stdout)",
    )
    parser.add_argument(
        "--model",
        default="base",
        choices=["tiny", "base", "small", "medium", "large-v3"],
        help="Whisper model size (default: base)",
    )
    parser.add_argument(
        "--language",
        default=None,
        help="Language code for transcription, e.g. 'en', 'fr', 'de' (default: auto-detect)",
    )
    parser.add_argument(
        "--task",
        default="transcribe",
        choices=["transcribe", "translate"],
        help="Task: 'transcribe' keeps original language, 'translate' translates to English",
    )
    parser.add_argument(
        "--format",
        default="txt",
        choices=["txt", "srt", "vtt", "docx"],
        help="Output format (default: txt). docx requires -o to specify output file.",
    )
    parser.add_argument(
        "--audio-dir",
        default=None,
        help="Directory to store downloaded audio (default: temp directory)",
    )
    parser.add_argument(
        "--keep-audio",
        action="store_true",
        help="Keep the downloaded audio file after transcription",
    )
    parser.add_argument(
        "--mode",
        default="simple",
        choices=["simple", "auto"],
        help=(
            "Transcription mode. 'simple': single-track transcription (legacy). "
            "'auto': multi-language pipeline — keeps EN/DE from original floor audio, "
            "uses DE interpreter track for other languages (default: simple)"
        ),
    )
    parser.add_argument(
        "--clean",
        action="store_true",
        help="Remove filler words and repetitions from the transcript",
    )
    parser.add_argument(
        "--audio-track",
        default=None,
        help=(
            "Audio track to download: 'or' for original floor, "
            "'de'/'en'/'fr'/... for interpreter channel (default: URL language)"
        ),
    )
    return parser.parse_args(argv)


def format_timestamp(seconds):
    """Convert seconds to HH:MM:SS,mmm format."""
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    millis = int((seconds % 1) * 1000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d},{millis:03d}"


def format_timestamp_vtt(seconds):
    """Convert seconds to HH:MM:SS.mmm format (VTT uses dot not comma)."""
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    millis = int((seconds % 1) * 1000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}.{millis:03d}"


def format_segments(segments, fmt):
    """Format transcription segments into the requested output format."""
    lines = []
    if fmt == "txt":
        for segment in segments:
            ts = format_timestamp(segment["start"]).split(",")[0]
            lang_tag = ""
            if "language" in segment:
                lang_tag = f" [{segment['language'].upper()}]"
            lines.append(f"[{ts}]{lang_tag} {segment['text'].strip()}")
    elif fmt == "srt":
        for i, segment in enumerate(segments, 1):
            start = format_timestamp(segment["start"])
            end = format_timestamp(segment["end"])
            lines.append(str(i))
            lines.append(f"{start} --> {end}")
            lines.append(segment["text"].strip())
            lines.append("")
    elif fmt == "vtt":
        lines.append("WEBVTT")
        lines.append("")
        for segment in segments:
            start = format_timestamp_vtt(segment["start"])
            end = format_timestamp_vtt(segment["end"])
            lines.append(f"{start} --> {end}")
            lines.append(segment["text"].strip())
            lines.append("")
    return "\n".join(lines)


def main(argv=None):
    args = parse_args(argv)

    if args.mode == "auto":
        _run_auto_mode(args)
    else:
        _run_simple_mode(args)


def _run_simple_mode(args):
    """Original single-track transcription."""
    try:
        print(f"Downloading audio from: {args.url}", file=sys.stderr)
        audio_path = download_audio(
            args.url, output_dir=args.audio_dir, audio_track=args.audio_track
        )
        print(f"Audio saved to: {audio_path}", file=sys.stderr)
    except Exception as e:
        print(f"Error downloading audio: {e}", file=sys.stderr)
        sys.exit(1)

    try:
        print(
            f"Transcribing with model '{args.model}' (language={args.language or 'auto'})...",
            file=sys.stderr,
        )
        segments = transcribe_audio(
            audio_path,
            model_size=args.model,
            language=args.language,
            task=args.task,
        )
    except Exception as e:
        print(f"Error during transcription: {e}", file=sys.stderr)
        sys.exit(1)

    if args.clean:
        segments = clean_segments(segments)

    _output_results(segments, args)

    if not args.keep_audio and not args.audio_dir:
        _cleanup_file(audio_path)


def _run_auto_mode(args):
    """Multi-language pipeline: original EN/DE + DE interpreter for other languages."""
    from video_transcriber.pipeline import run_pipeline

    try:
        segments = run_pipeline(
            args.url,
            model_size=args.model,
            output_dir=args.audio_dir,
        )
    except Exception as e:
        print(f"Error in auto pipeline: {e}", file=sys.stderr)
        sys.exit(1)

    _output_results(segments, args)


def _format_duration(seconds):
    """Format a duration in seconds as 'Xh Ym Zs', omitting zero components."""
    hours = int(seconds // 3600)
    minutes = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    parts = []
    if hours:
        parts.append(f"{hours}h")
    if minutes:
        parts.append(f"{minutes}m")
    if secs or not parts:
        parts.append(f"{secs}s")
    return " ".join(parts)


def write_docx(segments, path):
    """Write transcription segments to a Word (.docx) file."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.add_heading("Transcript", level=1)

    # Video duration from last segment
    if segments:
        duration = max(s["end"] for s in segments)
        p = doc.add_paragraph()
        run = p.add_run(f"Duration: {_format_duration(duration)}")
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(128, 128, 128)

    for segment in segments:
        text = segment["text"].strip()
        p = doc.add_paragraph()
        text_run = p.add_run(text)
        text_run.font.size = Pt(11)

    doc.save(path)


def _output_results(segments, args):
    """Format and output the transcription segments."""
    if args.format == "docx":
        if not args.output:
            print("Error: --format docx requires -o <file.docx>", file=sys.stderr)
            sys.exit(1)
        write_docx(segments, args.output)
        print(f"Transcript written to: {args.output}", file=sys.stderr)
        return

    output = format_segments(segments, args.format)

    if args.output:
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(output)
        print(f"Transcript written to: {args.output}", file=sys.stderr)
    else:
        print(output)


def _cleanup_file(path):
    """Remove a file, ignoring errors."""
    import os

    try:
        os.unlink(path)
    except OSError:
        pass


if __name__ == "__main__":
    main()
